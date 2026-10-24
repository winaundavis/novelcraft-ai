"""Services for importing markdown chapters and Word suggestion documents."""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from app.llm.audit import reset_llm_audit_context, set_llm_audit_context
from app.llm.provider import LLMConfigurationError, llm
from app.models.schemas import gen_id
from app.utils.word_count import count_words

try:
    import jieba.posseg as pseg

    JIEBA_AVAILABLE = True
except ImportError:  # pragma: no cover
    JIEBA_AVAILABLE = False


def _read_text_with_fallback(path: Path) -> str:
    # utf-8-sig 优先：能正确剥离 UTF-8 BOM；无 BOM 时 utf-8-sig 与 utf-8 行为一致
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError("unknown", b"", 0, 1, f"无法解码文件：{path.name}")


def _extract_chapter_number(filename: str) -> int:
    match = re.search(r"(\d+)", filename)
    return int(match.group(1)) if match else 9999


def _extract_title(content: str, filename: str) -> str:
    for line in content.splitlines()[:8]:
        stripped = line.strip()
        if stripped.startswith("# "):
            return stripped[2:].strip()
        if stripped.startswith("## "):
            return stripped[3:].strip()
    base = Path(filename).stem
    base = re.sub(r"^\d+[\s._-]*", "", base)
    return base or "未命名章节"


def _extract_characters(text: str) -> list[dict[str, Any]]:
    counter: dict[str, int] = {}
    stopwords = {
        "我们", "他们", "你们", "这里", "那里", "如果", "但是", "因为", "所以", "然后",
        "已经", "这个", "那个", "没有", "不是", "自己", "有人", "可以", "需要", "当前",
        "本章", "章节", "项目", "内容", "故事", "时候", "之后", "之前",
    }

    if JIEBA_AVAILABLE:
        for word, flag in pseg.cut(text):
            if flag in ("nr", "nrt") and 2 <= len(word) <= 4 and word not in stopwords:
                counter[word] = counter.get(word, 0) + 1
    else:
        for word in re.findall(r"(?<![\u4e00-\u9fff])([\u4e00-\u9fff]{2,4})(?![\u4e00-\u9fff])", text):
            if word not in stopwords:
                counter[word] = counter.get(word, 0) + 1

    results: list[dict[str, Any]] = []
    for name, count in sorted(counter.items(), key=lambda item: item[1], reverse=True):
        if count < 2:
            continue
        results.append({"name": name, "frequency": count, "aliases": []})
        if len(results) >= 20:
            break
    return results


def _extract_locations(text: str) -> list[dict[str, Any]]:
    suffixes = ("城", "港", "桥", "站", "街", "区", "湾", "岛", "楼", "巷", "馆", "寺", "宫")
    counter: dict[str, int] = {}
    pattern = re.compile(rf"([\u4e00-\u9fff]{{2,8}}(?:{'|'.join(suffixes)}))")

    for match in pattern.finditer(text):
        location = match.group(1)
        counter[location] = counter.get(location, 0) + 1

    return [
        {"name": name, "frequency": count}
        for name, count in sorted(counter.items(), key=lambda item: item[1], reverse=True)[:15]
        if count >= 2
    ]


def _extract_timeline(text: str, chapters: list[dict[str, Any]]) -> list[dict[str, Any]]:
    patterns = [
        (r"\d+小时后", "relative"),
        (r"第\d+天", "absolute"),
        (r"\d{4}年\d{1,2}月\d{1,2}日", "date"),
        (r"\d{4}年", "year"),
        (r"(清晨|早晨|中午|傍晚|深夜|凌晨)", "time"),
    ]

    chapter_offsets: list[tuple[int, int]] = []
    running = 0
    for index, chapter in enumerate(chapters, start=1):
        running += len(chapter.get("content", ""))
        chapter_offsets.append((index, running))

    events: list[dict[str, Any]] = []
    for pattern, kind in patterns:
        for match in re.finditer(pattern, text):
            position = match.start()
            chapter_number = 1
            for index, boundary in chapter_offsets:
                if position <= boundary:
                    chapter_number = index
                    break
            events.append(
                {
                    "text": match.group(0),
                    "type": kind,
                    "chapter": chapter_number,
                    "position": position,
                }
            )

    return sorted(events, key=lambda item: item["position"])[:30]


def _detect_foreshadowing(text: str) -> list[dict[str, Any]]:
    patterns = [
        (r"(秘密|真相|隐藏|异常|不对劲)", "隐藏线索"),
        (r"(似乎|仿佛).{0,10}(知道|明白|发现)", "暗示认知"),
        (r"(信号|笔记|档案|录音|坐标)", "物证线索"),
    ]

    findings: list[dict[str, Any]] = []
    for pattern, category in patterns:
        for match in re.finditer(pattern, text):
            start = max(0, match.start() - 24)
            end = min(len(text), match.end() + 40)
            findings.append(
                {
                    "text": match.group(0),
                    "category": category,
                    "context": text[start:end].replace("\n", " "),
                    "position": match.start(),
                }
            )

    return sorted(findings, key=lambda item: item["position"])[:20]


def build_analysis_summary(chapters: list[dict[str, Any]], analysis: dict[str, Any]) -> str:
    """剧情概览摘要文案（唯一来源）。

    扫描时生成一次；作者在导入前舍弃误识别的条目后要**再生成一次**，
    否则摘要里的人数会和实际保留的条目对不上。
    """
    return (
        f"已分析 {len(chapters)} 章，识别出 {len(analysis.get('characters') or [])} 个角色、"
        f"{len(analysis.get('locations') or [])} 个地点、"
        f"{len(analysis.get('timeline') or [])} 个时间节点。"
    )


def _analyze_full_text(text: str, chapters: list[dict[str, Any]]) -> dict[str, Any]:
    characters = _extract_characters(text)
    locations = _extract_locations(text)
    timeline = _extract_timeline(text, chapters)
    foreshadowing = _detect_foreshadowing(text)
    analysis = {
        "characters": characters,
        "locations": locations,
        "timeline": timeline,
        "foreshadowing": foreshadowing,
    }
    analysis["summary"] = build_analysis_summary(chapters, analysis)
    return analysis


async def scan_markdown_folder(folder_path: str) -> dict[str, Any]:
    folder = Path(folder_path)
    if not folder.is_dir():
        return {"success": False, "error": "文件夹路径不存在或不是目录。"}

    markdown_files = sorted(
        [path for path in folder.rglob("*.md") if path.is_file()],
        key=lambda path: (_extract_chapter_number(path.name), str(path.relative_to(folder))),
    )
    if not markdown_files:
        return {
            "success": False,
            "error": "该文件夹中未找到 Markdown 章节文件。",
            "folderPath": str(folder),
        }

    chapters: list[dict[str, Any]] = []
    full_text_parts: list[str] = []

    for order, file_path in enumerate(markdown_files, start=1):
        try:
            content = _read_text_with_fallback(file_path)
        except UnicodeDecodeError:
            return {
                "success": False,
                "error": f"内容解析失败，请检查文件编码：{file_path.name}",
                "folderPath": str(folder),
            }

        title = _extract_title(content, file_path.name)
        word_count = count_words(content)
        chapters.append(
            {
                "id": gen_id(),
                "fileName": file_path.name,
                "relativePath": file_path.relative_to(folder).as_posix(),
                "fullPath": str(file_path),
                "title": title,
                "content": content,
                "wordCount": word_count,
                "order": order,
            }
        )
        full_text_parts.append(content)

    full_text = "\n\n".join(full_text_parts)
    return {
        "success": True,
        "folderPath": str(folder),
        "totalFiles": len(chapters),
        "totalWords": sum(chapter["wordCount"] for chapter in chapters),
        "chapters": chapters,
        "analysis": _analyze_full_text(full_text, chapters),
    }


async def parse_suggestion_docx(file_path: str) -> dict[str, Any]:
    path = Path(file_path)
    if not path.is_file():
        return {"success": False, "error": "建议书文件不存在。"}

    try:
        from docx import Document
    except ImportError:
        return {"success": False, "error": "需要安装 python-docx 才能解析 .docx 文件。"}

    try:
        document = Document(str(path))
    except Exception:  # pragma: no cover
        # 不泄露服务器绝对路径（Codex P2：破损 docx 曾返回 Package not found at 'F:\...'）
        return {"success": False, "error": "Word 建议书解析失败：文件可能已损坏或格式不受支持，请检查文件后重试。"}

    paragraphs = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            paragraphs.append(
                {
                    "text": text,
                    "style": paragraph.style.name if paragraph.style else "Normal",
                }
            )

    tables = []
    for table in document.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append(rows)

    full_text = "\n".join(item["text"] for item in paragraphs)
    suggestions = await _parse_suggestions_with_ai(full_text)
    return {
        "success": True,
        "fileName": path.name,
        "paragraphs": paragraphs,
        "tables": tables,
        "suggestions": suggestions,
        "totalSuggestions": len(suggestions),
    }


async def _parse_suggestions_with_ai(text: str) -> list[dict[str, Any]]:
    if len(text.strip()) < 40:
        return _rule_based_parse_suggestions(text)

    prompt = f"""
请从下面这份小说修改建议书中提取可执行任务，输出 JSON 数组。每项字段包括：
- action: add / modify / delete / enhance / compress / restructure / style / character
- target: 目标章节、人物或情节
- description: 对建议的简洁描述
- priority: high / medium / low
- dependencies: 依赖的其他建议 id 列表，未知可空数组

建议书内容：
{text[:4000]}
"""

    token = set_llm_audit_context(role="word")
    try:
        reply = await llm.chat(
            [
                {"role": "system", "content": "你是专业编辑，请把建议书拆解成结构化执行项。"},
                {"role": "user", "content": prompt},
            ]
        )

        start = reply.find("[")
        end = reply.rfind("]") + 1
        if start >= 0 and end > start:
            raw_items = json.loads(reply[start:end])
            suggestions = []
            for index, item in enumerate(raw_items, start=1):
                suggestions.append(
                    {
                        "id": gen_id(),
                        "order": index,
                        "status": "pending",
                        "action": item.get("action", "modify"),
                        "target": item.get("target", "整体文本"),
                        "description": item.get("description", "").strip() or "待处理建议",
                        "priority": item.get("priority", "medium"),
                        "dependencies": item.get("dependencies", []),
                    }
                )
            if suggestions:
                return suggestions
    except LLMConfigurationError:
        return _rule_based_parse_suggestions(text)
    except Exception:
        pass
    finally:
        reset_llm_audit_context(token)

    return _rule_based_parse_suggestions(text)


def _rule_based_parse_suggestions(text: str) -> list[dict[str, Any]]:
    lines = [line.strip(" -\t") for line in text.splitlines() if line.strip()]
    suggestions: list[dict[str, Any]] = []

    keyword_rules = [
        ("新增", "add", "high"),
        ("增加", "add", "high"),
        ("补一场", "add", "high"),
        ("修改", "modify", "medium"),
        ("调整", "modify", "medium"),
        ("删除", "delete", "medium"),
        ("删减", "delete", "medium"),
        ("增强", "enhance", "high"),
        ("强化", "enhance", "high"),
        ("压缩", "compress", "medium"),
        ("重构", "restructure", "high"),
        ("风格", "style", "low"),
        ("人物", "character", "high"),
    ]

    for line in lines:
        action = "modify"
        priority = "medium"
        for keyword, mapped_action, mapped_priority in keyword_rules:
            if keyword in line:
                action = mapped_action
                priority = mapped_priority
                break
        suggestions.append(
            {
                "id": gen_id(),
                "order": len(suggestions) + 1,
                "status": "pending",
                "action": action,
                "target": line[:24] or "整体文本",
                "description": line[:120],
                "priority": priority,
                "dependencies": [],
            }
        )
        if len(suggestions) >= 12:
            break

    if not suggestions:
        suggestions.append(
            {
                "id": gen_id(),
                "order": 1,
                "status": "pending",
                "action": "enhance",
                "target": "整体文本",
                "description": "建议增强当前文本的情节推进、人物动机和描写层次。",
                "priority": "medium",
                "dependencies": [],
            }
        )

    return suggestions
